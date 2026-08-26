import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pypdf import PdfWriter
from docx import Document

from core.extract import Line, extract


class ExtractTests(unittest.TestCase):
    def test_table_based_docx_preserves_both_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table-resume.docx"
            document = Document()
            document.add_paragraph("Table Candidate")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Technical Skills\nCloud: AWS"
            table.cell(0, 1).text = "Education\nB.Tech | Example University | 2022"
            document.save(path)
            lines = extract(str(path))
        text = "\n".join(line.text for line in lines)
        self.assertIn("Cloud: AWS", text)
        self.assertIn("Example University", text)

    @patch("core.extract._from_scanned_pdf")
    def test_image_only_pdf_uses_ocr_fallback(self, ocr):
        ocr.return_value = [Line("OCR Candidate", bold=True), Line("SUMMARY", bold=True)]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with path.open("wb") as stream:
                writer.write(stream)
            lines = extract(str(path))
        self.assertEqual(lines[0].text, "OCR Candidate")
        ocr.assert_called_once()


if __name__ == "__main__":
    unittest.main()
